"""Real fixture pipeline, fake PDF transport captures the deck sent for export."""
from pathlib import Path
import json, tempfile
import pytest
from pptx import Presentation
from counter_risk.config import load_config
from counter_risk.pipeline.run import run_pipeline_with_config
from counter_risk.outputs.pdf_export import PDFExportGenerator

art = Path(__file__).parent
seen = []
config = load_config(Path('config/fixture_replay.yml')).model_copy(update={'export_pdf': True, 'include_concentration_table_in_ppt': True})
def capture_export(self, *, context):
    deck = Presentation(self.source_pptx)
    seen.append({'path': str(self.source_pptx), 'slides_at_export': len(deck.slides)})
    return ()  # No fabricated PDF, only a transport observation.
with tempfile.TemporaryDirectory(prefix='cr-final-deck-', dir=art) as tmp, pytest.MonkeyPatch.context() as m:
    m.setattr(PDFExportGenerator, 'generate', capture_export)
    out = run_pipeline_with_config(config, config_dir=Path('config'), output_dir=Path(tmp) / 'run')
    assert len(seen) == 1, seen
    final = len(Presentation(seen[0]['path']).slides)
    result = {'export_observation': seen, 'final_ppt_slides': final, 'difference': final-seen[0]['slides_at_export'], 'limitation': 'PDF rendering replaced with capture at generator boundary; actual pipeline and PPTX append executed.'}
    assert final == seen[0]['slides_at_export']+1, result
    (art/'lead-final-deck-proof.json').write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))
